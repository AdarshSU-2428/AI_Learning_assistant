from docx import Document

from core.loaders.loader import BaseLoader
from core.metadata import create_document_metadata


class DOCXLoader(BaseLoader):
    def extract_text(self):
        self.logger.info(f"Extracting text from {self.file_path.name}")

        try:
            document = Document(self.file_path)

            subject, content_type, source_type = self.infer_metadata_attributes()

            text = ""

            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"

            for table in document.tables:
                text += "\n"
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        row_text.append(cell.text.strip())

                    text += " | ".join(row_text) + "\n"

            self.logger.info(f"Successfully extracted text from DOCX (subject: {subject}, content_type: {content_type}, source_type: {source_type}).")

            meta = create_document_metadata(
                file_name=self.file_path.name,
                file_path=str(self.file_path),
                subject=subject,
                content_type=content_type,
                source_type=source_type
            )

            return [
                {
                    "text": text,
                    "metadata": meta,
                }
            ]

        except Exception as e:
            self.logger.error(f"Error while loading DOCX : {e}")
            raise