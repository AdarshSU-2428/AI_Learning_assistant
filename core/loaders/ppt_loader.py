from pptx import Presentation

from core.loaders.loader import BaseLoader
from config.settings import PDF_SLIDE_WINDOW_SIZE


class PPTLoader(BaseLoader):
    def __init__(self, file_path=None, window_size=PDF_SLIDE_WINDOW_SIZE):
        super().__init__(file_path)
        self.window_size = window_size

    def extract_text(self):
        self.logger.info(f"Extracting text from {self.file_path.name}")

        try:
            presentation = Presentation(self.file_path)
            raw_slides = []

            for slide_num, slide in enumerate(presentation.slides, start=1):
                slide_text = ""
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text += shape.text + "\n"

                if slide_text and slide_text.strip():
                    raw_slides.append((slide_num, slide_text.strip()))

            units = []
            if not raw_slides:
                return units

            if len(raw_slides) > 1:
                for i in range(len(raw_slides)):
                    window = raw_slides[i : i + self.window_size]
                    combined_text = "\n\n".join(
                        f"[Slide {s}]:\n{txt}" for s, txt in window
                    )
                    start_s = window[0][0]
                    end_s = window[-1][0]
                    slide_meta = (
                        start_s if start_s == end_s else f"{start_s}-{end_s}"
                    )

                    units.append(
                        {
                            "text": combined_text,
                            "metadata": {"slide": slide_meta},
                        }
                    )
            else:
                for slide_num, slide_text in raw_slides:
                    units.append(
                        {
                            "text": slide_text,
                            "metadata": {"slide": slide_num},
                        }
                    )

            self.logger.info(
                f"Successfully extracted {len(units)} units from {len(presentation.slides)} slides."
            )

            return units

        except Exception as e:
            self.logger.error(f"Error while reading PPT: {e}")
            raise